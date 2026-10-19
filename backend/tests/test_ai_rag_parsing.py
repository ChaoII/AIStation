"""解析测试：文本解码、docx/xlsx 表格、PDF 页偏移、MinerU 结果转换、分派。"""
import asyncio
import io


def test_decode_text_handles_gbk():
    from app.plugin.module_ai.knowledge.parsing import text as tparse

    assert tparse.decode_text("中文".encode("gbk")) == "中文"


def test_csv_to_markdown():
    from app.plugin.module_ai.knowledge.parsing import text as tparse

    md = tparse.csv_to_markdown("名称,数量\n螺丝,10\n")
    assert "| 名称 | 数量 |" in md and "| 螺丝 | 10 |" in md


def test_docx_tables_become_markdown():
    from docx import Document

    from app.plugin.module_ai.knowledge.parsing import office

    document = Document()
    document.add_paragraph("段落一")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "列A"
    table.cell(0, 1).text = "列B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    buf = io.BytesIO()
    document.save(buf)
    body, meta = office.parse_docx(buf.getvalue())
    assert "段落一" in body
    assert "| 列A | 列B |" in body
    assert "| 1 | 2 |" in body
    assert meta["tables"] == 1


def test_xlsx_multi_sheet():
    from openpyxl import Workbook

    from app.plugin.module_ai.knowledge.parsing import office

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "库存"
    sheet.append(["名称", "数量"])
    sheet.append(["螺丝", 10])
    workbook.create_sheet("备注").append(["说明"])
    buf = io.BytesIO()
    workbook.save(buf)
    body, meta = office.parse_xlsx(buf.getvalue())
    assert "## 库存" in body and "| 螺丝 | 10 |" in body
    assert "## 备注" in body
    assert meta["sheets"] == 2


def test_pdf_page_offsets_and_scanned_detection():
    from pypdf import PdfWriter

    from app.plugin.module_ai.knowledge.parsing import pdf as pdf_parser

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    body, offsets, _meta = pdf_parser.parse_pdf(buf.getvalue())
    assert offsets[0] == 0 and len(offsets) == 2
    assert pdf_parser.is_scanned(body) is True
    assert pdf_parser.is_scanned("这是一段足够长的可提取文本" * 5) is False


def test_html_table_conversion_and_page_offsets():
    from app.plugin.module_ai.knowledge.parsing import mineru

    md = (
        "前文<table><tr><th>列A</th><th>列B</th></tr>"
        "<tr><td>1</td><td>2</td></tr></table>后文"
    )
    out = mineru.html_tables_to_markdown(md)
    assert "| 列A | 列B |" in out and "| 1 | 2 |" in out
    assert "<table" not in out

    text = "第一页内容\n第二页内容"
    offsets = mineru.page_offsets_from_content_list(
        [
            {"page_idx": 0, "text": "第一页内容"},
            {"page_idx": 1, "text": "第二页内容"},
        ],
        text,
    )
    assert offsets == [0, len("第一页内容\n")]


def test_parse_document_dispatch():
    from app.plugin.module_ai.knowledge.parsing import parse_document

    doc = asyncio.run(parse_document("a.txt", "纯文本内容".encode("utf-8")))
    assert "纯文本内容" in doc.text and doc.meta["parser"] == "text"

    doc2 = asyncio.run(parse_document("a.unknown", b"fallback bytes"))
    assert doc2.meta["parser"] == "text"


def test_parse_document_falls_back_on_broken_docx():
    from app.plugin.module_ai.knowledge.parsing import parse_document

    doc = asyncio.run(parse_document("broken.docx", b"not a real docx"))
    assert doc.meta.get("fallback") is True
    assert "not a real docx" in doc.text


def test_web_import_creates_task(test_client, auth_headers, monkeypatch):
    """网页导入：抓取成功后建文档并入队（不依赖真实网络）。"""
    import uuid

    from app.plugin.module_ai.knowledge.parsing import web

    async def fake_parse_url(url, kb_id=None):
        return "网页正文内容", [0], {"parser": "web", "url": url}

    monkeypatch.setattr(web, "parse_url", fake_parse_url)
    kb_id = test_client.post(
        "/api/v1/ai/knowledge/bases/create",
        json={"name": f"pytest_web_{uuid.uuid4().hex[:8]}", "enabled": True},
        headers=auth_headers,
    ).json()["data"]["id"]
    try:
        r = test_client.post(
            f"/api/v1/ai/knowledge/bases/{kb_id}/web-import",
            json={"url": "http://example.com/a"},
            headers=auth_headers,
        )
        assert r.status_code == 200, r.text
        data = r.json()["data"]
        assert data["task_id"] and data["id"]

        missing = test_client.post(
            f"/api/v1/ai/knowledge/bases/{kb_id}/web-import",
            json={"url": ""},
            headers=auth_headers,
        )
        assert missing.status_code == 400
        assert missing.json()["success"] is False
    finally:
        test_client.request(
            "DELETE",
            "/api/v1/ai/knowledge/bases/delete",
            json=[kb_id],
            headers=auth_headers,
        )
